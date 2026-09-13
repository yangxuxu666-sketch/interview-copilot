from __future__ import annotations

import asyncio
from contextlib import asynccontextmanager, suppress
from copy import deepcopy
import hmac
import importlib.util
import io
import json
import logging
from pathlib import Path
import platform
import secrets as security
import sys
import time
from typing import Literal
from urllib.parse import urlparse
from uuid import uuid4
import zipfile

from fastapi import FastAPI, File, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field, field_validator

from .storage import Store, now
from . import intelligence
from .overlay import OverlayProcess

STATIC = Path(__file__).parent / "static"
MAX_IMPORT = 8 * 1024 * 1024


def desktop_capabilities():
    """Inspect installed components without opening audio or requesting access."""
    if sys.platform == "darwin":
        try:
            supported = int(platform.mac_ver()[0].split(".")[0]) >= 13
        except (ValueError, IndexError):
            supported = False
        return {"platform": "macOS", "audio": supported and all(
            importlib.util.find_spec(name) is not None
            for name in ("objc", "ScreenCaptureKit", "CoreAudio", "CoreMedia", "dispatch")),
            "audio_note": "Mac 首次监听需要允许系统的屏幕与系统音频录制权限；本工具只处理系统播放声音，不启用麦克风。"}
    if sys.platform == "win32":
        return {"platform": "Windows", "audio": importlib.util.find_spec("pyaudiowpatch") is not None,
                "audio_note": ""}
    return {"platform": platform.system(), "audio": False, "audio_note": "当前系统可手动输入问题，暂不支持声音监听。"}


class SettingsPatch(BaseModel):
    model_config = ConfigDict(extra="ignore")
    deepseek_model: str | None = Field(None, min_length=1, max_length=100, pattern=r"^[\w.\-]+$")
    deepseek_key: str | None = Field(None, max_length=512)
    cloud_asr_key: str | None = Field(None, max_length=512)
    delete_deepseek_key: bool | None = None
    delete_cloud_asr_key: bool | None = None
    asr_provider: Literal["local", "cloud", "qwen"] | None = None
    whisper_model: Literal["tiny", "base", "small"] | None = None
    asr_language: Literal["zh", "en", "auto"] | None = None
    asr_quality: Literal["fast", "balanced", "accurate"] | None = None
    cloud_asr_url: str | None = Field(None, max_length=2048)
    cloud_asr_model: str | None = Field(None, max_length=100)
    silence_ms: int | None = Field(None, ge=350, le=2000)
    min_speech_ms: int | None = Field(None, ge=150, le=1500)
    max_segment_s: float | None = Field(None, ge=4, le=30)
    energy_threshold: float | None = Field(None, ge=0.001, le=0.1)
    auto_answer: bool | None = None
    answer_language: Literal["zh", "en", "auto"] | None = None
    answer_style: Literal["concise", "star", "technical"] | None = None
    device_id: int | None = Field(None, ge=0)

    @field_validator("cloud_asr_url")
    @classmethod
    def https_endpoint(cls, value):
        if value:
            p = urlparse(value)
            if p.scheme != "https" or not p.hostname or p.username or p.password or p.query or p.fragment:
                raise ValueError("请输入不含用户名、密码和查询参数的 HTTPS 转写接口。")
        return value

    @field_validator("deepseek_key", "cloud_asr_key")
    @classmethod
    def key_is_header_safe(cls, value):
        if value and (not value.isascii() or any(ord(c) < 32 for c in value)):
            raise ValueError("密钥格式不正确")
        return value


class ProfilePatch(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str | None = None
    name: str | None = Field(None, max_length=100)
    company: str | None = Field(None, max_length=200)
    role: str | None = Field(None, max_length=200)
    resume: str | None = Field(None, max_length=50000)
    jd: str | None = Field(None, max_length=50000)
    company_notes: str | None = Field(None, max_length=50000)
    materials: str | None = Field(None, max_length=50000)
    source_notes: str | None = Field(None, max_length=10000)
    custom_prompt: str | None = Field(None, max_length=4000)


class IdInput(BaseModel):
    id: str


class OverlayInput(BaseModel):
    answer_id: str | None = None


class AnswerInput(BaseModel):
    question: str = Field("", max_length=8000)
    mode: Literal["answer", "prepare", "review"] = "answer"


class AudioInput(BaseModel):
    device_id: int | None = Field(None, ge=0)


class FeedbackInput(BaseModel):
    answer_id: str
    rating: Literal["useful", "improve"]
    note: str = Field("", max_length=2000)


def new_session(profile):
    return {"id": uuid4().hex, "at": now(), "name": profile.get("name", "面试记录"),
            "profile_id": profile.get("id"), "listening": False, "transcripts": [], "answers": []}


class Controller:
    def __init__(self, store: Store):
        from .audio import AudioService
        self.store = store
        self.loop = asyncio.get_running_loop()
        self.clients: dict[WebSocket, asyncio.Queue] = {}
        self.audio = AudioService(store.root, self.thread_event, self.thread_transcript)
        self.session = store.load_session() or new_session(store.active_profile())
        if self.session.get("profile_id") != store.state["active_profile_id"]:
            if self.session.get("transcripts") or self.session.get("answers"):
                store.save_session(self.session, archive=True)
            self.session = new_session(store.active_profile())
            store.save_session(self.session)
        self.answer_task = None
        self.current_answer = None
        self.auto_task = None
        self.auto_fragments = []
        self.last_auto_question = ""
        self.disconnect_task = None
        self.answer_lock = asyncio.Lock()
        self.control_lock = asyncio.Lock()
        self.closing = False
        self.components = {}
        self.audio_epoch = 0
        self.overlay = OverlayProcess()
        self.overlay_selected_id = None
        self.overlay_revision = 0
        self.overlay_show_revision = 0
        self.overlay_changed = asyncio.Event()

    def touch_overlay(self):
        previous = self.overlay_changed
        self.overlay_changed = asyncio.Event()
        self.overlay_revision += 1
        previous.set()

    def overlay_snapshot(self):
        answers = self.session["answers"]
        selected = next((a for a in answers if a["id"] == self.overlay_selected_id), None)
        selected = selected or (answers[-1] if answers else None)
        answer = {k: selected.get(k) for k in ("id", "question", "text", "status", "error")} if selected else None
        return {"revision": self.overlay_revision, "show_revision": self.overlay_show_revision,
                "answer": answer, "listening": self.session["listening"],
                "audio_message": self.components.get("asr", {}).get("message", "等待生成回答")}

    def snapshot(self):
        audio_status = self.audio.status()
        return {"settings": self.store.public_settings(), "profiles": deepcopy(self.store.state["profiles"]),
                "active_profile_id": self.store.state["active_profile_id"], "session": deepcopy(self.session),
                "capabilities": {**desktop_capabilities(),
                                 "cloud_edition": bool(getattr(sys, "_interview_cloud_edition", False)),
                                 "local_asr": importlib.util.find_spec("faster_whisper") is not None,
                                 "audio_status": audio_status, "secret_error": self.store.secret_error},
                "components": deepcopy(self.components)}

    async def publish(self, event):
        if event.get("type") == "answer_start":
            self.overlay_selected_id = event.get("id")
        if event.get("type") == "status":
            self.components[event.get("component", "audio")] = event
            if event.get("component") == "audio":
                if event.get("state") in ("error", "idle"):
                    self.session["listening"] = False
                elif event.get("state") == "listening":
                    self.session["listening"] = True
        if event.get("type") != "level":
            self.touch_overlay()
        for ws, queue in list(self.clients.items()):
            if queue.full():
                if event.get("type") == "level":
                    continue
                self.clients.pop(ws, None)
                asyncio.create_task(ws.close(code=1013))
                continue
            queue.put_nowait(event)

    def thread_event(self, event):
        if not self.closing and not self.loop.is_closed():
            epoch = self.audio_epoch
            self.loop.call_soon_threadsafe(lambda: asyncio.create_task(self.publish(event)) if epoch == self.audio_epoch else None)

    def thread_transcript(self, event):
        if not self.closing and not self.loop.is_closed():
            epoch = self.audio_epoch
            self.loop.call_soon_threadsafe(lambda: asyncio.create_task(self.transcript(event)) if epoch == self.audio_epoch else None)

    async def transcript(self, event):
        if event.get("endpoint") and self.session["listening"] and self.auto_fragments and self.store.state["settings"]["auto_answer"]:
            await self.cancel_auto(clear=False)
            self.auto_task = asyncio.create_task(self.auto_answer())
            return
        if not self.session["listening"] or not event.get("text", "").strip():
            return
        item = {"id": uuid4().hex, "at": now(), **event}
        self.session["transcripts"].append(item)
        self.store.save_session(self.session)
        await self.publish({"type": "transcript", **item})
        if self.store.state["settings"]["auto_answer"]:
            self.auto_fragments.append(item["text"])
            self.auto_fragments = self.auto_fragments[-8:]
            await self.cancel_auto(clear=False)
            self.auto_task = asyncio.create_task(self.auto_answer(continuation=item.get("continuation", False)))

    async def auto_answer(self, continuation=False):
        try:
            await self.publish({"type": "auto_pending", "text": " ".join(self.auto_fragments)})
            # Forced long-speech slices normally merge with the next endpoint.
            # A bounded fallback handles speech ending exactly on the slice boundary.
            delay = self.store.state["settings"]["max_segment_s"] * 2 + 5 if continuation else 0.35
            await asyncio.sleep(delay)
            question = " ".join(self.auto_fragments).strip()[-8000:]
            if not self.session["listening"] or not self.store.state["settings"]["auto_answer"]:
                return
            if len(question) < 4 or question == self.last_auto_question:
                self.auto_fragments.clear()
                return
            self.auto_fragments.clear()
            self.last_auto_question = question
            await self.start_answer(question, "answer", automatic=True)
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            await self.publish({"type": "warning", "message": safe_error(exc)})

    async def cancel_auto(self, clear=True):
        task = self.auto_task
        self.auto_task = None
        if task and task is not asyncio.current_task() and not task.done():
            task.cancel()
            with suppress(asyncio.CancelledError):
                await task
        if clear:
            self.auto_fragments.clear()

    async def stop_answer(self):
        async with self.answer_lock:
            await self._cancel_answer_task()

    async def _cancel_answer_task(self):
        if self.answer_task and not self.answer_task.done():
            self.answer_task.cancel()
            with suppress(asyncio.CancelledError):
                await self.answer_task
        # Cancelling a task before its first execution bypasses its finally block.
        if self.current_answer and self.current_answer.get("status") == "streaming":
            self.current_answer.update(status="cancelled", total_ms=0)
            self.store.save_session(self.session)
            await self.publish({"type": "answer_done", **self.current_answer, "cancelled": True})
        self.answer_task = None
        self.current_answer = None

    async def start_answer(self, question, mode, automatic=False):
        async with self.control_lock:
            return await self._start_answer(question, mode, automatic)

    async def _start_answer(self, question, mode, automatic=False):
        if not self.store.secrets.get("deepseek_key"):
            raise HTTPException(400, "请先在设置中填写 DeepSeek API Key。")
        question = question.strip()
        if not question:
            if mode == "answer":
                raise HTTPException(400, "请先输入或选择面试问题。")
            question = "请结合当前简历、岗位和公司资料生成面试准备清单。" if mode == "prepare" else "请复盘本场面试问题与回答，结合我的反馈列出改进和练习建议。"
        if not automatic:
            await self.cancel_auto()
        async with self.answer_lock:
            await self._cancel_answer_task()
            answer = {"id": uuid4().hex, "question": question, "mode": mode, "text": "", "at": now(),
                      "status": "streaming", "first_token_ms": None, "total_ms": None, "sources": []}
            profile = self.store.active_profile()
            settings = dict(self.store.state["settings"])
            history = deepcopy(self.session["answers"][-12:])
            self.session["answers"].append(answer)
            self.store.save_session(self.session)
            await self.publish({"type": "answer_start", **answer})
            self.current_answer = answer
            self.answer_task = asyncio.create_task(self.generate(answer, profile, settings, history))
            return {"id": answer["id"]}

    async def generate(self, answer, profile, settings, history):
        started = time.perf_counter()
        await self.publish({"type": "status", "component": "llm", "state": "loading", "message": "正在组织回答"})
        try:
            messages, sources = intelligence.build_messages(profile, answer["question"], history, settings, answer["mode"])
            answer["sources"] = sources
            async for delta in intelligence.stream_answer(settings, dict(self.store.secrets), messages):
                if not delta:
                    continue
                if answer["first_token_ms"] is None:
                    answer["first_token_ms"] = round((time.perf_counter() - started) * 1000)
                answer["text"] += delta
                await self.publish({"type": "answer_delta", "id": answer["id"], "text": delta,
                                    "first_token_ms": answer["first_token_ms"]})
            answer["status"] = "done"
        except asyncio.CancelledError:
            answer["status"] = "cancelled"
            raise
        except Exception as exc:
            answer["status"] = "error"
            answer["error"] = safe_error(exc)
            await self.publish({"type": "answer_error", "id": answer["id"], "message": answer["error"]})
        finally:
            answer["total_ms"] = round((time.perf_counter() - started) * 1000)
            self.store.save_session(self.session)
            await self.publish({"type": "answer_done", **answer, "cancelled": answer["status"] == "cancelled"})
            await self.publish({"type": "status", "component": "llm", "state": "idle", "message": "等待下一个问题"})

    async def stop_audio(self):
        self.session["listening"] = False
        self.audio_epoch += 1
        await self.cancel_auto()
        await asyncio.to_thread(self.audio.stop)
        self.store.save_session(self.session)

    async def reset(self):
        await self.stop_audio()
        await self.stop_answer()
        if self.session["transcripts"] or self.session["answers"]:
            self.store.save_session(self.session, archive=True)
        self.session = new_session(self.store.active_profile())
        self.last_auto_question = ""
        self.store.save_session(self.session)
        await self.publish({"type": "session_reset", "session": deepcopy(self.session)})

    async def disconnected(self):
        await asyncio.sleep(15)
        # The answer window can be the remaining active view of an interview.
        while not self.clients and self.overlay.running and not self.closing:
            await asyncio.sleep(2)
        if not self.clients:
            await self.stop_audio()
            await self.stop_answer()

    async def close(self):
        self.closing = True
        self.touch_overlay()
        if self.disconnect_task:
            self.disconnect_task.cancel()
        await self.stop_audio()
        await self.stop_answer()
        await asyncio.to_thread(self.audio.close)
        await asyncio.to_thread(self.overlay.close)


def safe_error(exc):
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    if isinstance(exc, (RuntimeError, ValueError)):
        message = str(exc)
        # Third-party exception text never travels to UI unless it is a curated module error.
        if type(exc).__module__.startswith("interview_copilot"):
            return message[:600]
    return "操作未完成，请检查服务设置、网络或音频设备后重试。"


def extract_document(filename, content):
    suffix = Path(filename).suffix.lower()
    warning = None
    try:
        if suffix in (".txt", ".md"):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise HTTPException(400, "Word 文档解压后过大，请导出为纯文本。")
            doc = Document(io.BytesIO(content))
            blocks = [p.text for p in doc.paragraphs]
            blocks.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
            text = "\n".join(blocks)
        elif suffix == ".pdf":
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(content))
            if pdf.is_encrypted:
                raise HTTPException(400, "PDF 已加密，请先解密或粘贴文字。")
            if len(pdf.pages) > 100:
                raise HTTPException(400, "PDF 超过 100 页，请导入与面试相关的部分。")
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        else:
            raise HTTPException(400, "支持 TXT、MD、DOCX 和带文字的 PDF 文件。")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "无法读取文档，请检查格式或复制文字导入。") from None
    text = text.replace("\x00", "").strip()
    if not text:
        raise HTTPException(400, "没有提取到文字。扫描版 PDF 需要先进行 OCR，或直接粘贴内容。")
    if len(text) > 50000:
        text = text[:50000]
        warning = "文件文字超过 50,000 字符，已截取前 50,000 字符，请检查后保存。"
    return {"filename": Path(filename).name, "text": text, "warning": warning}


def export_markdown(session):
    lines = [f"# {session.get('name', '面试记录')}", "", f"开始时间：{session.get('at', '')}", "", "## 转写", ""]
    for t in session["transcripts"]:
        lines.extend([f"- [{t.get('at', '')}] {t['text']}"])
    lines.extend(["", "## 问题与辅助回答", ""])
    for a in session["answers"]:
        lines.extend([f"### {a['question']}", "", a["text"] or "（未生成内容）", "",
                      f"状态：{a.get('status', '')}；首字：{a.get('first_token_ms')} ms；总耗时：{a.get('total_ms')} ms", ""])
        if a.get("feedback"):
            lines.extend([f"反馈：{a['feedback']['rating']} / {a['feedback'].get('note', '')}", ""])
    return "\n".join(lines)


def create_app(data_dir: Path, token: str | None = None, port: int = 8765):
    token = token or security.token_urlsafe(32)
    origins = {f"http://127.0.0.1:{port}", f"http://localhost:{port}"}
    hosts = {f"127.0.0.1:{port}", f"localhost:{port}"}

    @asynccontextmanager
    async def lifespan(app):
        app.state.controller = Controller(Store(data_dir))
        yield
        await app.state.controller.close()

    app = FastAPI(title="面试伴航", lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
    app.state.token = token
    app.state.shutdown_callback = None

    @app.middleware("http")
    async def local_boundary(request: Request, call_next):
        if request.headers.get("host") not in hosts:
            return JSONResponse({"detail": "仅允许本机访问。"}, status_code=403)
        if request.url.path.startswith("/api/"):
            if not hmac.compare_digest(request.cookies.get("interview_session", ""), token):
                return JSONResponse({"detail": "请通过桌面启动器打开应用。"}, status_code=401)
            origin = request.headers.get("origin")
            if origin and origin not in origins:
                return JSONResponse({"detail": "拒绝跨站请求。"}, status_code=403)
            if request.method not in ("GET", "HEAD", "OPTIONS") and origin not in origins:
                return JSONResponse({"detail": "需要来自本机应用的请求。"}, status_code=403)
            if int(request.headers.get("content-length", "0") or "0") > MAX_IMPORT + 1024 * 64:
                return JSONResponse({"detail": "文件过大，单个文件上限为 8 MB。"}, status_code=413)
        response = await call_next(request)
        response.headers["Cache-Control"] = "no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; object-src 'none'"
        return response

    @app.exception_handler(RequestValidationError)
    async def bad_input(request, exc):
        fields = ", ".join(str(e["loc"][-1]) for e in exc.errors())
        return JSONResponse({"detail": f"输入格式或范围不正确：{fields}"}, status_code=422)

    @app.exception_handler(KeyError)
    async def missing(request, exc):
        return JSONResponse({"detail": str(exc).strip("'")}, status_code=404)

    @app.exception_handler(ValueError)
    async def invalid(request, exc):
        return JSONResponse({"detail": str(exc)}, status_code=400)

    @app.get("/health")
    async def health():
        return {"application": "interview-companion", "version": "0.1.0"}

    @app.get("/launch")
    async def launch(token: str = ""):
        if not hmac.compare_digest(token, app.state.token):
            raise HTTPException(401, "启动链接已失效，请重新运行启动器。")
        response = RedirectResponse("/", status_code=303)
        response.set_cookie("interview_session", token, httponly=True, samesite="strict", path="/")
        return response

    @app.get("/")
    async def index():
        return FileResponse(STATIC / "index.html")

    @app.get("/api/state")
    async def state():
        return app.state.controller.snapshot()

    @app.post("/api/settings")
    async def settings(patch: SettingsPatch):
        c = app.state.controller
        data = patch.model_dump(exclude_unset=True)
        data = {k: v for k, v in data.items() if v is not None or k == "device_id"}
        if getattr(sys, "_interview_cloud_edition", False) and data.get("asr_provider") == "local":
            raise HTTPException(400, "此分享版支持云端语音识别，请选择阿里云千问或其他云端服务。")
        capture_fields = {"device_id", "asr_provider", "whisper_model", "asr_language", "asr_quality", "cloud_asr_url", "cloud_asr_model", "cloud_asr_key", "delete_cloud_asr_key", "silence_ms", "min_speech_ms", "max_segment_s", "energy_threshold"}
        if c.session["listening"] and any(k in capture_fields and v != c.store.state["settings"].get(k) for k, v in data.items()):
            raise HTTPException(409, "请先暂停监听，再修改声音或转写设置。")
        try:
            result = c.store.update_settings(data)
        except RuntimeError as exc:
            raise HTTPException(400, str(exc)) from None
        if not result["auto_answer"]:
            await c.cancel_auto()
        return result

    @app.post("/api/settings/test")
    async def test_settings():
        c = app.state.controller
        try:
            return await intelligence.test_connection(c.store.state["settings"], c.store.secrets)
        except Exception as exc:
            raise HTTPException(400, safe_error(exc)) from None

    @app.get("/api/audio/devices")
    async def devices():
        try:
            return {"devices": await asyncio.to_thread(app.state.controller.audio.devices)}
        except Exception as exc:
            return {"devices": [], "error": safe_error(exc)}

    @app.get("/api/audio/status")
    async def audio_status():
        return app.state.controller.audio.status()

    @app.post("/api/audio/start")
    async def start_audio(body: AudioInput):
        c = app.state.controller
        async with c.control_lock:
            if c.session["listening"]:
                return {"listening": True}
            try:
                c.audio_epoch += 1
                await asyncio.to_thread(c.audio.configure, dict(c.store.state["settings"]), dict(c.store.secrets))
                await asyncio.to_thread(c.audio.start, body.device_id)
                audio_status = c.audio.status()
                c.session["listening"] = bool(audio_status.get("listening"))
                if not c.session["listening"]:
                    raise HTTPException(400, audio_status.get("message", "音频启动失败，请检查设备与模型。"))
            except Exception as exc:
                c.session["listening"] = False
                raise HTTPException(400, safe_error(exc)) from None
        return {"listening": True}

    @app.post("/api/audio/stop")
    async def stop_audio():
        c = app.state.controller
        async with c.control_lock:
            await c.stop_audio()
        return {"listening": False}

    @app.post("/api/audio/warmup")
    async def warmup():
        c = app.state.controller
        async with c.control_lock:
            if c.session["listening"]:
                raise HTTPException(409, "请先暂停监听。")
            try:
                await asyncio.to_thread(c.audio.configure, dict(c.store.state["settings"]), dict(c.store.secrets))
                await asyncio.to_thread(c.audio.warmup)
            except Exception as exc:
                raise HTTPException(400, safe_error(exc)) from None
        return {"status": "loading", "message": "正在准备本地语音模型；首次使用需要下载。"}

    @app.post("/api/profiles")
    async def save_profile(body: ProfilePatch):
        return app.state.controller.store.save_profile(body.model_dump(exclude_none=True))

    @app.post("/api/profiles/active")
    async def active_profile(body: IdInput):
        c = app.state.controller
        async with c.control_lock:
            if c.session["listening"]:
                raise HTTPException(409, "请先暂停监听，再切换面试档案。")
            if body.id != c.store.state["active_profile_id"]:
                # Validate first, archive the old session, then bind the new profile.
                if not any(p["id"] == body.id for p in c.store.state["profiles"]):
                    raise KeyError("资料档案不存在")
                c.store.activate(body.id)
                await c.reset()
        return c.snapshot()

    @app.delete("/api/profiles/{profile_id}")
    async def delete_profile(profile_id: str):
        c = app.state.controller
        async with c.control_lock:
            if c.session["listening"]:
                raise HTTPException(409, "请先暂停监听，再删除档案。")
            was_active = c.store.state["active_profile_id"] == profile_id
            c.store.delete_profile(profile_id)
            if was_active:
                await c.reset()
        return c.snapshot()

    @app.post("/api/import")
    async def import_file(file: UploadFile = File(...)):
        try:
            content = await file.read(MAX_IMPORT + 1)
            if len(content) > MAX_IMPORT:
                raise HTTPException(413, "单个文件不能超过 8 MB。")
            return await asyncio.to_thread(extract_document, file.filename or "file.txt", content)
        finally:
            await file.close()

    @app.post("/api/answer")
    async def answer(body: AnswerInput):
        return await app.state.controller.start_answer(body.question, body.mode)

    @app.post("/api/answer/stop")
    async def stop_answer():
        c = app.state.controller
        await c.cancel_auto()
        await c.stop_answer()
        return {"stopped": True}

    @app.post("/api/overlay/open")
    async def open_overlay(body: OverlayInput | None = None):
        c = app.state.controller
        answer_id = body.answer_id if body else None
        if answer_id and not any(a["id"] == answer_id for a in c.session["answers"]):
            raise HTTPException(404, "本场回答不存在。")
        c.overlay_selected_id = answer_id
        try:
            await asyncio.to_thread(c.overlay.open, port, token, c.store.root)
        except RuntimeError as exc:
            raise HTTPException(400, safe_error(exc)) from None
        c.overlay_show_revision += 1
        c.touch_overlay()
        return {"opened": True}

    @app.post("/api/overlay/close")
    async def close_overlay():
        await asyncio.to_thread(app.state.controller.overlay.close)
        return {"closed": True}

    @app.post("/api/overlay/select")
    async def select_overlay(body: IdInput):
        c = app.state.controller
        if not any(a["id"] == body.id for a in c.session["answers"]):
            raise HTTPException(404, "本场回答不存在。")
        c.overlay_selected_id = body.id
        c.touch_overlay()
        return {"selected": body.id}

    @app.get("/api/overlay/state")
    async def overlay_state(after: int = -1):
        c = app.state.controller
        if after == c.overlay_revision and not c.closing:
            try:
                await asyncio.wait_for(c.overlay_changed.wait(), timeout=12)
            except asyncio.TimeoutError:
                pass
        return c.overlay_snapshot()

    @app.post("/api/session/new")
    async def reset_session():
        c = app.state.controller
        async with c.control_lock:
            await c.reset()
        return {"session": deepcopy(c.session)}

    @app.get("/api/session/export")
    async def export():
        return Response(export_markdown(app.state.controller.session), media_type="text/markdown; charset=utf-8",
                        headers={"Content-Disposition": 'attachment; filename="interview-session.md"'})

    @app.post("/api/feedback")
    async def feedback(body: FeedbackInput):
        c = app.state.controller
        answer = next((a for a in c.session["answers"] if a["id"] == body.answer_id), None)
        if answer is None:
            raise KeyError("回答不存在")
        answer["feedback"] = {"rating": body.rating, "note": body.note, "at": now()}
        c.store.save_session(c.session)
        return {"saved": True}

    @app.get("/api/sessions")
    async def sessions():
        return {"sessions": app.state.controller.store.list_sessions()}

    @app.get("/api/sessions/{session_id}")
    async def session(session_id: str):
        return app.state.controller.store.get_session(session_id)

    @app.post("/api/shutdown")
    async def shutdown():
        await app.state.controller.close()
        if app.state.shutdown_callback:
            asyncio.get_running_loop().call_later(0.4, app.state.shutdown_callback)
        return {"stopped": True}

    @app.websocket("/ws")
    async def socket(ws: WebSocket):
        if (ws.headers.get("host") not in hosts or ws.headers.get("origin") not in origins or
                not hmac.compare_digest(ws.cookies.get("interview_session", ""), token)):
            await ws.close(code=1008)
            return
        c = app.state.controller
        await ws.accept()
        queue = asyncio.Queue(maxsize=512)
        c.clients[ws] = queue
        if c.disconnect_task:
            c.disconnect_task.cancel()
        await ws.send_json({"type": "snapshot", **c.snapshot()})

        async def sender():
            while True:
                event = await queue.get()
                await ws.send_json(event)

        send_task = asyncio.create_task(sender())
        try:
            while True:
                message = await ws.receive_text()
                if message == "ping":
                    queue.put_nowait({"type": "pong"})
        except (WebSocketDisconnect, RuntimeError):
            pass
        finally:
            send_task.cancel()
            with suppress(asyncio.CancelledError, RuntimeError):
                await send_task
            c.clients.pop(ws, None)
            if not c.clients and not c.closing:
                c.disconnect_task = asyncio.create_task(c.disconnected())

    app.mount("/static", StaticFiles(directory=STATIC), name="static")
    return app
