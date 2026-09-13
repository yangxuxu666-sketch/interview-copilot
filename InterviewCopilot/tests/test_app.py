import asyncio
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from interview_copilot.app import create_app, extract_document
from interview_copilot.storage import Store, _dpapi
from interview_copilot.intelligence import IntelligenceError

BASE = "http://127.0.0.1:8765"


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(Path(self.temp.name), "test-only-token")
        self.client = TestClient(self.app, base_url=BASE, headers={"Origin": BASE})
        self.client.__enter__()
        self.client.get("/launch?token=test-only-token")

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.temp.cleanup()

    def test_private_api_and_websocket_require_local_session(self):
        self.client.cookies.clear()
        self.assertEqual(self.client.get("/api/state").status_code, 401)
        self.assertEqual(self.client.get("/launch?token=wrong").status_code, 401)
        with self.assertRaises(Exception):
            with self.client.websocket_connect("ws://127.0.0.1:8765/ws"):
                pass

    def test_local_cookie_and_cross_origin_rejection(self):
        r = self.client.get("/launch?token=test-only-token", follow_redirects=False)
        self.assertIn("HttpOnly", r.headers["set-cookie"])
        self.assertIn("SameSite=strict", r.headers["set-cookie"])
        self.assertEqual(self.client.post("/api/session/new", headers={"Origin": "https://evil.example"}).status_code, 403)
        self.assertEqual(self.client.get("/api/state", headers={"Host": "evil.example"}).status_code, 403)
        self.assertEqual(self.client.get("/api/state").status_code, 200)

    def test_default_is_stopped_bilingual_and_no_secret(self):
        body = self.client.get("/api/state").json()
        self.assertFalse(body["session"]["listening"])
        self.assertEqual(body["settings"]["answer_language"], "auto")
        self.assertEqual(body["settings"]["asr_language"], "auto")
        self.assertNotIn("deepseek_key", body["settings"])

    def test_validation_never_echoes_secret_and_rejects_non_tls(self):
        r = self.client.post("/api/settings", json={"deepseek_key": "SECRET\nVALUE"})
        self.assertEqual(r.status_code, 422)
        self.assertNotIn("SECRET", r.text)
        self.assertEqual(self.client.post("/api/settings", json={"cloud_asr_url": "http://example.com/asr"}).status_code, 422)
        self.assertEqual(self.client.post("/api/settings", json={"silence_ms": 1}).status_code, 422)

    def test_profiles_persist_and_active_switch_starts_new_session(self):
        old = self.client.get("/api/state").json()
        profile = self.client.post("/api/profiles", json={"name": "准备B", "company": "示例公司", "resume": "真实履历"}).json()
        state = self.client.post("/api/profiles/active", json={"id": profile["id"]}).json()
        self.assertEqual(state["active_profile_id"], profile["id"])
        self.assertNotEqual(old["session"]["id"], state["session"]["id"])
        disk = Store(Path(self.temp.name))
        self.assertEqual(disk.active_profile()["resume"], "真实履历")
        self.assertEqual(self.client.delete("/api/profiles/not-real").status_code, 404)
        self.client.delete("/api/profiles/" + old["active_profile_id"])
        self.assertEqual(self.client.delete("/api/profiles/" + profile["id"]).status_code, 400)

    def test_import_handles_unicode_invalid_file_and_docx_tables(self):
        r = self.client.post("/api/import", files={"file": ("简历.txt", "团队管理经历".encode("utf-8"), "text/plain")})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["text"], "团队管理经历")
        self.assertEqual(self.client.post("/api/import", files={"file": ("bad.exe", b"MZ")}).status_code, 400)
        from docx import Document
        doc = Document()
        doc.add_paragraph("项目负责人")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "简历表格证据"
        data = io.BytesIO()
        doc.save(data)
        result = extract_document("resume.docx", data.getvalue())
        self.assertIn("简历表格证据", result["text"])
        self.assertEqual(self.client.post("/api/import", files={"file": ("empty.pdf", b"broken")}).status_code, 400)

    def test_connection_errors_are_actionable(self):
        async def failure(*args):
            raise IntelligenceError("API Key 无效，请重新填写。")
        with patch("interview_copilot.app.intelligence.test_connection", failure):
            r = self.client.post("/api/settings/test")
        self.assertEqual(r.status_code, 400)
        self.assertIn("API Key", r.json()["detail"])

    def test_manual_answer_requires_key_without_network(self):
        r = self.client.post("/api/answer", json={"question": "请介绍自己"})
        self.assertEqual(r.status_code, 400)

    def test_stream_feedback_archive_and_export(self):
        async def fake_stream(*args):
            yield "先说结论。"
            await asyncio.sleep(0.01)
            yield "再说明真实经历。"
        self.app.state.controller.store.secrets["deepseek_key"] = "local-test-placeholder"
        with patch("interview_copilot.app.intelligence.stream_answer", fake_stream):
            with self.client.websocket_connect("ws://127.0.0.1:8765/ws", headers={"origin": BASE}) as ws:
                self.assertEqual(ws.receive_json()["type"], "snapshot")
                result = self.client.post("/api/answer", json={"question": "How do you lead a team?"}).json()
                events = []
                for _ in range(15):
                    event = ws.receive_json()
                    events.append(event)
                    if event["type"] == "answer_done":
                        break
                final = events[-1]
                self.assertEqual(final["status"], "done")
                self.assertEqual(final["text"], "先说结论。再说明真实经历。")
                self.assertIsNotNone(final["first_token_ms"])
                self.assertTrue(any(e["type"] == "answer_delta" for e in events))
        self.assertEqual(self.client.post("/api/feedback", json={"answer_id": result["id"], "rating": "improve", "note": "更简短"}).status_code, 200)
        export = self.client.get("/api/session/export")
        self.assertIn("更简短", export.text)
        old_id = self.client.get("/api/state").json()["session"]["id"]
        self.client.post("/api/session/new")
        archived = self.client.get("/api/sessions").json()["sessions"]
        self.assertEqual(archived[0]["id"], old_id)
        self.assertEqual(self.client.get("/api/sessions/" + old_id).json()["answers"][0]["feedback"]["note"], "更简短")

    def test_cancel_preserves_partial_answer(self):
        async def waiting(*args):
            yield "partial"
            await asyncio.sleep(30)
        self.app.state.controller.store.secrets["deepseek_key"] = "local-test-placeholder"
        with patch("interview_copilot.app.intelligence.stream_answer", waiting):
            with self.client.websocket_connect("ws://127.0.0.1:8765/ws", headers={"origin": BASE}) as ws:
                ws.receive_json()
                self.client.post("/api/answer", json={"question": "Q?"})
                for _ in range(10):
                    if ws.receive_json()["type"] == "answer_delta":
                        break
                self.client.post("/api/answer/stop")
                state = self.client.get("/api/state").json()
                self.assertEqual(state["session"]["answers"][-1]["status"], "cancelled")
                self.assertEqual(state["session"]["answers"][-1]["text"], "partial")

    @unittest.skipUnless(os.name == "nt", "Windows DPAPI")
    def test_dpapi_encrypted_key_roundtrip_and_blank_preserves(self):
        key = "sk-unit-test-not-a-real-key"
        self.assertEqual(_dpapi(_dpapi(key.encode()), decrypt=True), key.encode())
        self.assertEqual(self.client.post("/api/settings", json={"deepseek_key": key}).status_code, 200)
        self.client.post("/api/settings", json={"deepseek_key": ""})
        loaded = Store(Path(self.temp.name))
        self.assertEqual(loaded.secrets["deepseek_key"], key)
        self.assertNotIn(key, (Path(self.temp.name) / "state.json").read_text(encoding="utf-8"))
        self.assertNotIn(key.encode(), (Path(self.temp.name) / "secrets.dpapi").read_bytes())
        self.client.post("/api/settings", json={"delete_deepseek_key": True})
        self.assertFalse(self.client.get("/api/state").json()["settings"]["deepseek_key_set"])


if __name__ == "__main__":
    unittest.main()
